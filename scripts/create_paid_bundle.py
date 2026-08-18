from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

OUTPUT = Path("/home/ubuntu/webdev-static-assets")
OUTPUT.mkdir(parents=True, exist_ok=True)
WORKBOOK_PATH = OUTPUT / "FlipProfit_Deal_Screening_Workbook.xlsx"
CHECKLIST_PATH = OUTPUT / "FlipProfit_Field_Checklist.pdf"
ZIP_PATH = OUTPUT / "FlipProfit_Deal_Screening_Bundle.zip"

CARBON = "1B1D1E"
ORANGE = "FF5C35"
CREAM = "F7F6F0"
PAPER = "EFEDE5"
GREEN = "A9E681"
AMBER = "F4C46A"
RED = "FF775F"
GRAY = "69716D"
THIN = Side(style="thin", color="C9C8C0")


def title(ws, text, subtitle):
    ws.sheet_view.showGridLines = False
    ws.merge_cells("B2:H2")
    ws["B2"] = text
    ws["B2"].font = Font(name="Arial", size=22, bold=True, color="FFFFFF")
    ws["B2"].fill = PatternFill("solid", fgColor=CARBON)
    ws["B2"].alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 33
    ws.merge_cells("B3:H3")
    ws["B3"] = subtitle
    ws["B3"].font = Font(name="Arial", size=10, color=GRAY, italic=True)
    ws["B3"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[3].height = 28
    ws.column_dimensions["A"].width = 3
    for column, width in {"B": 28, "C": 18, "D": 4, "E": 27, "F": 20, "G": 4, "H": 18}.items():
        ws.column_dimensions[column].width = width


def section(ws, row, text, end_column="H"):
    ws.merge_cells(f"B{row}:{end_column}{row}")
    cell = ws[f"B{row}"]
    cell.value = text
    cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=CARBON)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 20


def set_input(cell):
    cell.fill = PatternFill("solid", fgColor="FFF8D9")
    cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    cell.alignment = Alignment(vertical="center")


def set_output(cell):
    cell.fill = PatternFill("solid", fgColor=CREAM)
    cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    cell.font = Font(name="Arial", bold=True, color=CARBON)
    cell.alignment = Alignment(vertical="center")


def build_workbook():
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"
    title(overview, "FLIPPROFIT", "Deal Screening Bundle — Workbook Edition")
    section(overview, 5, "START HERE")
    steps = [
        "1. Open the Work Order sheet. Enter every value you know. Use 0 only when you have confirmed there is no cost.",
        "2. Leave nothing blank if you want a decision. A blank is treated as unknown, never as zero.",
        "3. Check the decision trace, then copy the completed work order to the Ledger only after you have verified your numbers.",
        "4. Keep the PDF Field Checklist with you during inspection, transport, title, and sale preparation.",
    ]
    for index, text in enumerate(steps, start=6):
        overview.merge_cells(f"B{index}:H{index}")
        overview[f"B{index}"] = text
        overview[f"B{index}"].alignment = Alignment(wrap_text=True, vertical="center")
        overview[f"B{index}"].font = Font(name="Arial", size=10, color=CARBON)
        overview.row_dimensions[index].height = 28
    section(overview, 12, "ACCOUNTABILITY RULES")
    rules = [
        ("No seeded deals", "This workbook begins blank. It contains no sample vehicles, prices, claims, or sales results."),
        ("Every number has a place", "Use the notes column to record where the number came from: auction bill, repair quote, title office, or direct confirmation."),
        ("The workbook does not appraise", "It calculates the values you enter. It does not predict sale price, condition, demand, or profit."),
        ("Keep a record", "Use the Ledger sheet to retain confirmed deals and export it when you need a copy outside Excel."),
    ]
    for row, (label, explanation) in enumerate(rules, start=13):
        overview[f"B{row}"] = label
        overview[f"B{row}"].font = Font(name="Arial", size=10, bold=True, color=ORANGE)
        overview.merge_cells(f"C{row}:H{row}")
        overview[f"C{row}"] = explanation
        overview[f"C{row}"].alignment = Alignment(wrap_text=True, vertical="center")
        overview.row_dimensions[row].height = 31
    overview.merge_cells("B19:H19")
    overview["B19"] = "Yellow cells are yours to enter. Gray cells calculate automatically once all required fields are complete."
    overview["B19"].fill = PatternFill("solid", fgColor="FFF8D9")
    overview["B19"].font = Font(name="Arial", size=10, italic=True, color=CARBON)
    overview["B19"].alignment = Alignment(wrap_text=True, vertical="center")
    overview.row_dimensions[19].height = 28

    work = workbook.create_sheet("Work Order")
    title(work, "DEAL SCREENING WORK ORDER", "Enter confirmed values in yellow cells. Blank fields stop the workbook from making a decision.")
    section(work, 5, "DEAL FACTS")
    work["B6"] = "Work order name"
    set_input(work["C6"])
    work["B7"] = "Expected sale price"
    work["B8"] = "Buy price"
    work["B9"] = "Auction / purchase fee"
    work["B10"] = "Transport / tow"
    work["B11"] = "Repair labor"
    work["B12"] = "Parts"
    work["B13"] = "Title / document costs"
    work["B14"] = "Detail / cleanup"
    work["B15"] = "Other costs"
    work["B16"] = "Selling fee rate"
    work["B17"] = "Fixed selling fee"
    work["B18"] = "Advertising / listing"
    work["B19"] = "Target profit"
    for row in range(7, 20):
        set_input(work[f"C{row}"])
        work[f"C{row}"].number_format = "0.0%" if row == 16 else "$#,##0.00"
        work[f"B{row}"].font = Font(name="Arial", size=10, color=CARBON)
        work.row_dimensions[row].height = 21
    work["B21"] = "Notes / evidence source"
    work.merge_cells("C21:H23")
    set_input(work["C21"])
    work["C21"].alignment = Alignment(wrap_text=True, vertical="top")
    for row in range(21, 24):
        work.row_dimensions[row].height = 20
    section(work, 25, "DECISION TRACE")
    trace = [
        ("Completeness", '=IF(COUNTBLANK(C7:C19)>0,"NEEDS REVIEW","READY")', "General"),
        ("Acquisition cost", '=IF(F26<>"READY","",SUM(C8:C15))', "$#,##0.00"),
        ("Selling fee", '=IF(F26<>"READY","",C7*C16+C17)', "$#,##0.00"),
        ("Selling costs", '=IF(F26<>"READY","",F28+C18)', "$#,##0.00"),
        ("Net profit", '=IF(F26<>"READY","",C7-F27-F29)', "$#,##0.00"),
        ("Return on cash", '=IF(OR(F26<>"READY",F27=0),"",F30/F27)', "0.0%"),
        ("Break-even sale price", '=IF(OR(F26<>"READY",1-C16=0),"",(F27+C17+C18)/(1-C16))', "$#,##0.00"),
        ("Sale price for target", '=IF(OR(F26<>"READY",1-C16=0),"",(F27+C17+C18+C19)/(1-C16))', "$#,##0.00"),
        ("Decision", '=IF(F26<>"READY","NEEDS REVIEW",IF(F30>=C19,"BUY",IF(F30>=0,"BORDERLINE","PASS")))', "General"),
    ]
    for row, (label, formula, number_format) in enumerate(trace, start=26):
        work[f"E{row}"] = label
        work[f"E{row}"].font = Font(name="Arial", size=10, color=CARBON)
        work[f"F{row}"] = formula
        work[f"F{row}"].number_format = number_format
        set_output(work[f"F{row}"])
        work.row_dimensions[row].height = 22
    work["E35"] = "Formula"
    work.merge_cells("F35:H36")
    work["F35"] = "Net profit = expected sale price − acquisition cost − selling costs. The target price accounts for the percentage fee rate entered above."
    work["F35"].alignment = Alignment(wrap_text=True, vertical="center")
    work["F35"].font = Font(name="Arial", size=9, italic=True, color=GRAY)
    work.conditional_formatting.add("F34", CellIsRule(operator="equal", formula=['"BUY"'], fill=PatternFill("solid", fgColor=GREEN)))
    work.conditional_formatting.add("F34", CellIsRule(operator="equal", formula=['"BORDERLINE"'], fill=PatternFill("solid", fgColor=AMBER)))
    work.conditional_formatting.add("F34", CellIsRule(operator="equal", formula=['"PASS"'], fill=PatternFill("solid", fgColor=RED)))
    work.freeze_panes = "B6"

    ledger = workbook.create_sheet("Ledger")
    title(ledger, "CONFIRMED DEAL LEDGER", "Copy only confirmed work orders here. This sheet starts blank and contains no example entries.")
    headers = ["Date", "Work order", "Expected sale", "Buy price", "Acquisition cost", "Selling costs", "Net profit", "ROI", "Decision", "Evidence / notes"]
    for col, header in enumerate(headers, start=2):
        cell = ledger.cell(row=5, column=col, value=header)
        cell.fill = PatternFill("solid", fgColor=CARBON)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ledger.row_dimensions[5].height = 28
    for row in range(6, 106):
        for col in range(2, 12):
            cell = ledger.cell(row=row, column=col)
            cell.border = Border(bottom=Side(style="hair", color="D8D6CE"))
            cell.alignment = Alignment(vertical="center", wrap_text=(col == 11))
        for col in [4, 5, 6, 7, 8]:
            ledger.cell(row=row, column=col).number_format = "$#,##0.00"
        ledger.cell(row=row, column=9).number_format = "0.0%"
    for column, width in {"B": 14, "C": 25, "D": 15, "E": 15, "F": 17, "G": 15, "H": 15, "I": 12, "J": 14, "K": 36}.items():
        ledger.column_dimensions[column].width = width
    ledger.freeze_panes = "B6"
    ledger.auto_filter.ref = "B5:K105"

    checklist = workbook.create_sheet("Checklist")
    title(checklist, "FIELD CHECKLIST", "Use this before you buy, during pickup, and before you list. Check only items you have actually verified.")
    categories = {
        "Before you bid / buy": ["Confirm the vehicle or equipment ID and title status.", "Write down the total buy price limit from your completed Work Order.", "Confirm every auction, payment, or buyer fee.", "Record the expected sale-price evidence source."],
        "Condition and repair": ["Photograph visible condition issues.", "List required parts and obtain a repair estimate.", "Identify work you can verify versus work that is still unknown.", "Enter confirmed labor, parts, cleanup, and storage costs."],
        "Transport and paperwork": ["Confirm transport or tow cost.", "Confirm title, registration, document, or release fees.", "Confirm the payment method and any processing fee.", "Record where all key documents are stored."],
        "Before you list": ["Enter an explicit $0 for any confirmed zero cost.", "Confirm marketplace fee assumptions and their current category.", "Read the Decision Trace; resolve every blank before relying on status.", "Save or export the final record."],
    }
    row = 5
    for category, items in categories.items():
        section(checklist, row, category)
        row += 1
        for item in items:
            checklist[f"B{row}"] = "☐"
            checklist[f"C{row}"] = item
            checklist[f"C{row}"].alignment = Alignment(wrap_text=True, vertical="center")
            checklist[f"D{row}"] = "Evidence / note"
            set_input(checklist[f"E{row}"])
            checklist.merge_cells(start_row=row, start_column=5, end_row=row, end_column=8)
            checklist.row_dimensions[row].height = 28
            row += 1
        row += 1
    checklist.column_dimensions["B"].width = 5
    checklist.column_dimensions["C"].width = 60
    checklist.column_dimensions["D"].width = 18
    checklist.column_dimensions["E"].width = 18
    checklist.freeze_panes = "B5"

    for sheet in workbook.worksheets:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_view.zoomScale = 95

    workbook.save(WORKBOOK_PATH)


class ChecklistPDF(FPDF):
    def header(self):
        self.set_fill_color(27, 29, 30)
        self.rect(0, 0, 210, 25, "F")
        self.set_xy(14, 7)
        self.set_font("Helvetica", "B", 20)
        self.set_text_color(255, 92, 53)
        self.cell(0, 8, "FLIPPROFIT", ln=1)
        self.set_x(14)
        self.set_font("Helvetica", "", 9)
        self.set_text_color(247, 246, 240)
        self.cell(0, 5, "FIELD CHECKLIST - BLANK BY DESIGN")

    def footer(self):
        self.set_y(-14)
        self.set_font("Helvetica", "", 8)
        self.set_text_color(105, 113, 109)
        self.cell(0, 6, "Use only verified facts. Blank is unknown, not zero.", align="L")
        self.cell(0, 6, f"Page {self.page_no()}", align="R")


def build_checklist_pdf():
    pdf = ChecklistPDF(format="letter")
    pdf.set_auto_page_break(auto=True, margin=19)
    pdf.add_page()
    pdf.set_y(34)
    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(60, 65, 64)
    pdf.multi_cell(0, 5, "This is a working inspection sheet for an actual deal. It contains no example values. Mark only what you have checked and record the source of each cost or condition note.")
    sections = [
        ("Before you bid / buy", ["Confirm vehicle or equipment ID and title status.", "Write the maximum buy price from your completed Work Order.", "Confirm all auction, buyer, payment, and release fees.", "Write the source used for expected sale price."]),
        ("Condition and repair", ["Photograph condition issues and document visible damage.", "List needed parts, labor, cleanup, and unknown repair work.", "Get a repair estimate or flag the cost as unknown.", "Record any storage, insurance, or time-sensitive cost."]),
        ("Transport and paperwork", ["Confirm transport or tow charge.", "Confirm title, tax, document, registration, or release cost.", "Store document and payment evidence where you can find it."]),
        ("Before you list", ["Enter an explicit $0 only for confirmed zero costs.", "Verify marketplace fee category and current rule.", "Resolve every blank field before relying on a BUY/PASS result.", "Save the completed work order and export your ledger."]),
    ]
    for heading, items in sections:
        if heading == "Before you list":
            pdf.add_page()
        pdf.ln(4)
        pdf.set_fill_color(27, 29, 30)
        pdf.set_text_color(255, 255, 255)
        pdf.set_font("Helvetica", "B", 11)
        pdf.cell(0, 7, heading, fill=True, ln=1)
        pdf.set_text_color(38, 42, 42)
        pdf.set_font("Helvetica", "", 10)
        for item in items:
            pdf.multi_cell(0, 7, f"[ ]  {item}\n    Evidence / note: _________________________________________________________________")
            pdf.ln(1)
    pdf.output(CHECKLIST_PATH)


def package_bundle():
    with ZipFile(ZIP_PATH, "w", ZIP_DEFLATED) as archive:
        archive.write(WORKBOOK_PATH, WORKBOOK_PATH.name)
        archive.write(CHECKLIST_PATH, CHECKLIST_PATH.name)


if __name__ == "__main__":
    build_workbook()
    build_checklist_pdf()
    package_bundle()
    print(f"Created {WORKBOOK_PATH}")
    print(f"Created {CHECKLIST_PATH}")
    print(f"Created {ZIP_PATH}")
